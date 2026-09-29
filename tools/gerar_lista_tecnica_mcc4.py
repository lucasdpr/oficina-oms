"""Gera JS/Paineis/listaTecnicaMCC4.js a partir de dados/lista_tecnica_mcc4.xlsx.

Rodar de novo sempre que a planilha for atualizada:
    python3 tools/gerar_lista_tecnica_mcc4.py
"""
import json
import re
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "dados" / "lista_tecnica_mcc4.xlsx"
SAIDA = RAIZ / "JS" / "Paineis" / "listaTecnicaMCC4.js"

# Ordem importa: a primeira regra que bater decide o conjunto.
REGRAS = [
    ("placaLarga", ["foot roll", "footroll", "placa larga", "backup"]),
    ("placaEstreita", ["placa estreita", "telesc", "caixa de transm", "caixas de transm", "cardan",
                       "sanfonad", "edge", "guia", "tartaruga", "calço placa", "chaveta"]),
    ("tubulao", ["tubulão"]),
    ("carcacaMovel", ["lado movel", "lado móvel"]),
    ("carcacaFixa", ["cilindro", "haste", "kit vedação", "lado fixo"]),
]


# Sugestões pros itens com APLICAÇÃO vazia na planilha — só onde dá pra
# ter segurança pelo próprio texto do item. Aparecem no painel marcadas
# "a confirmar". O que não está aqui fica sem peça até a planilha ser
# preenchida (a APLICAÇÃO da planilha sempre ganha da sugestão).
SUGESTOES = {
    # água principal
    "1775081": ("tubulao", "Água Principal"), "9442127": ("tubulao", "Água Principal"),
    "1210714": ("tubulao", "Engate Água Principal"), "1625816": ("tubulao", "Flexível água"),
    # hidráulica dos cilindros
    "8012769": ("carcacaFixa", "Engate Hidráulica"), "8012768": ("carcacaFixa", "Engate Hidráulica"),
    "9414084": ("carcacaFixa", "Filtro hidráulico"), "9412817": ("carcacaFixa", "Filtro hidráulico"),
    "9412818": ("carcacaFixa", "Elemento filtro hidráulico"), "9384085": ("carcacaFixa", "Elemento filtro hidráulico"),
    "9245787": ("carcacaFixa", "Elemento filtro hidráulico"), "1756583": ("carcacaFixa", "Tubulação Hidráulica 12mm"),
    "8221482": ("carcacaFixa", "Tubulação Hidráulica 12mm"), "8360095": ("carcacaFixa", "Tubulação Hidráulica 12mm"),
    "9376806": ("carcacaFixa", "Rótula do cilindro"), "8893209": ("carcacaFixa", "Tomador de Pressão"),
    # placa estreita / telescópio / régua
    "8012895": ("placaEstreita", "Gaxeta Telescópio"), "1059438": ("placaEstreita", "Porca de ajuste"),
    "8034239": ("placaEstreita", "Régua"), "8034280": ("placaEstreita", "Régua"),
}


def conjunto_de(aplicacao):
    a = (aplicacao or "").lower()
    if not a:
        return "geral"
    for chave, termos in REGRAS:
        if any(t in a for t in termos):
            return chave
    return "geral"


def bitola_de(texto):
    t = texto.upper()
    if not re.search(r"PARAF|ARRUELA|PORCA", t):
        return None
    m = re.search(r"\bM\s?(\d{1,2})\b", t)
    return f"M{m.group(1)}" if m else None


def num(v):
    if v is None or v == "":
        return None
    try:
        f = float(str(v).replace(",", "."))
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def main():
    ws = openpyxl.load_workbook(ENTRADA, data_only=True).active
    atualizado = ws.cell(1, 6).value
    itens = []
    for linha in ws.iter_rows(min_row=4, values_only=True):
        codigo, texto = linha[0], linha[1]
        if codigo is None or not texto:
            continue
        texto = str(texto).strip()
        aplicacao = (str(linha[7]).strip() if linha[7] else "")
        conjunto = conjunto_de(aplicacao)
        sugestao = False
        if not aplicacao and str(codigo).strip() in SUGESTOES:
            conjunto, aplicacao = SUGESTOES[str(codigo).strip()]
            sugestao = True
        elif not aplicacao:
            conjunto = None
        itens.append({
            "codigo": str(codigo).strip(),
            "texto": texto,
            "qtd": num(linha[2]),
            "un": linha[5] or "",
            "aplicacao": aplicacao,
            "conjunto": conjunto,
            "sugestao": sugestao,
            "bitola": bitola_de(texto),
        })
    dados = {"equipamento": "Molde MCC#4", "atualizado": str(atualizado or ""), "itens": itens}
    SAIDA.write_text(
        "// GERADO por tools/gerar_lista_tecnica_mcc4.py a partir de dados/lista_tecnica_mcc4.xlsx — não editar à mão.\n"
        "export const LISTA_TECNICA_MCC4 = " + json.dumps(dados, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )
    contagem = {}
    for i in itens:
        contagem[i["conjunto"]] = contagem.get(i["conjunto"], 0) + 1
    print(len(itens), "itens", contagem)


if __name__ == "__main__":
    main()
